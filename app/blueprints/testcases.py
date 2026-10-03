"""R03 生成 / R04 查看编辑 / R05 保存 / R06 查询 / R07 删除 / R08 导出。"""
import csv
import io
import json
from datetime import datetime

from flask import (Blueprint, flash, jsonify, redirect, render_template, request,
                   send_file, url_for)
from flask_login import current_user, login_required

from .. import llm
from ..extensions import db
from ..models import (CASE_STATUS_CONFIRMED, CASE_STATUS_DRAFT, PRIORITY_NAMES,
                      Requirement, TestCase)
from ..security import flash_errors, log_action, next_code

bp = Blueprint("testcases", __name__, url_prefix="/testcases")

PER_PAGE = 10


def _active_requirements():
    return (db.session.query(Requirement)
            .filter_by(is_deleted=False)
            .order_by(Requirement.updated_at.desc()).all())


def _requirements_map() -> dict:
    """供生成页前端预览需求内容使用的 需求ID -> 内容 映射。"""
    return {
        str(r.id): {
            "code": r.code,
            "name": r.name,
            "description": r.description or "",
            "business_rules": r.business_rules or "",
            "case_count": r.testcases.filter_by(is_deleted=False).count(),
        }
        for r in db.session.query(Requirement).filter_by(is_deleted=False).all()
    }


def _render_generate(**kwargs):
    return render_template(
        "testcases/generate.html",
        requirements=_active_requirements(),
        configured=llm.is_configured(),
        requirements_map=_requirements_map(),
        **kwargs,
    )


def _filters():
    return {
        "q": (request.args.get("q") or "").strip(),
        "req": request.args.get("req", type=int),
        "category": (request.args.get("category") or "").strip(),
        "priority": (request.args.get("priority") or "").strip(),
        "status": (request.args.get("status") or "").strip(),
    }


def _apply_filters(query, f):
    if f["q"]:
        like = f"%{f['q']}%"
        query = query.filter(db.or_(TestCase.title.like(like),
                                    TestCase.code.like(like),
                                    TestCase.expected_result.like(like)))
    if f["req"]:
        query = query.filter(TestCase.requirement_id == f["req"])
    if f["category"]:
        query = query.filter(TestCase.category == f["category"])
    if f["priority"]:
        query = query.filter(TestCase.priority == f["priority"])
    if f["status"]:
        query = query.filter(TestCase.status == f["status"])
    return query


# ---------------------------- R06 查询 ----------------------------
@bp.route("/")
@login_required
def list_testcases():
    f = _filters()
    page = request.args.get("page", 1, type=int)
    query = _apply_filters(db.session.query(TestCase).filter_by(is_deleted=False), f)
    pagination = query.order_by(TestCase.updated_at.desc()).paginate(
        page=page, per_page=PER_PAGE, error_out=False)
    categories = [c[0] for c in db.session.query(TestCase.category).distinct().all() if c[0]]
    return render_template("testcases/list.html", pagination=pagination, filters=f,
                           requirements=_active_requirements(), categories=categories,
                           priorities=PRIORITY_NAMES)


# ---------------------------- R03 自动生成 ----------------------------
def _wants_json() -> bool:
    """前端通过 AJAX 提交时返回 JSON，便于展示处理状态（4.1 性能需求）。"""
    return (request.headers.get("X-Requested-With") == "XMLHttpRequest"
            or request.is_json)


def _validate_requirement(item) -> list:
    """3.3 业务规则：经过基本校验的需求才能提交生成。"""
    errors = []
    if item is None or item.is_deleted:
        errors.append("请选择一个有效的软件需求。")
        return errors
    if not (item.description or "").strip():
        errors.append("该需求尚未填写【需求描述】，无法提交生成。")
    elif len(item.description.strip()) < 10:
        errors.append("【需求描述】内容过短（少于 10 个字符），请补充完整后再提交生成。")
    return errors


@bp.route("/generate", methods=["GET", "POST"])
@login_required
def generate():
    if request.method == "POST":
        req_id = request.form.get("requirement_id", type=int)
        count = request.form.get("count", 5, type=int)
        count = min(max(count or 5, 1), 10)
        extra = (request.form.get("extra") or "").strip()
        item = db.session.get(Requirement, req_id) if req_id else None

        errors = _validate_requirement(item)
        if errors:
            if _wants_json():
                return jsonify({"ok": False, "error": " ".join(errors)}), 400
            for e in errors:
                flash(e, "danger")
            return _render_generate(results=None, selected=item)

        description = item.description
        if extra:
            description = f"{description}\n\n【补充说明】\n{extra}"

        try:
            results = llm.generate_testcases(item.name, description,
                                             item.business_rules, count=count)
        except llm.LLMError as exc:
            log_action("生成测试用例", target_type="requirement", target_id=item.id,
                       detail=f"{item.code} 生成失败：{exc}", result="failed")
            if _wants_json():
                return jsonify({"ok": False, "error": str(exc)}), 502
            flash(str(exc), "danger")
            return _render_generate(selected=item, results=None)

        log_action("生成测试用例", target_type="requirement", target_id=item.id,
                   detail=f"{item.code} 调用模型生成 {len(results)} 条用例")

        if _wants_json():
            return jsonify({
                "ok": True,
                "count": len(results),
                "results": results,
                "requirement": {"id": item.id, "code": item.code, "name": item.name},
            })

        flash(f"模型已返回 {len(results)} 条用例，请核对后确认保存。", "success")
        return _render_generate(selected=item, results=results)

    # 支持从需求列表/详情页带参跳转，直接预选目标需求
    preselect = db.session.get(Requirement, request.args.get("requirement_id", type=int)) \
        if request.args.get("requirement_id", type=int) else None
    if preselect is not None and preselect.is_deleted:
        preselect = None

    return _render_generate(results=None, selected=preselect)


@bp.route("/save-batch", methods=["POST"])
@login_required
def save_batch():
    """保存经用户确认的生成结果（R05）。3.3 业务规则：生成结果必须由用户确认。"""
    req_id = request.form.get("requirement_id", type=int)
    item = db.session.get(Requirement, req_id) if req_id else None
    if item is None or item.is_deleted:
        flash("关联的软件需求无效，无法保存。", "danger")
        return redirect(url_for("testcases.generate"))

    # 前端可能删除任意一行，故以 row_total 为准逐个探测，避免索引断档导致漏存
    total = request.form.get("row_total", type=int)
    if not total or total < 1:
        total = sum(1 for k in request.form if k.endswith("-title"))
    total = min(total, 50)

    saved, skipped = 0, 0
    for idx in range(total):
        title_key = f"rows-{idx}-title"
        if title_key not in request.form:
            continue
        if request.form.get(f"rows-{idx}-selected") != "on":
            continue
        title = (request.form.get(title_key) or "").strip()
        if not title:
            skipped += 1
            continue

        case = TestCase(
            code=next_code(TestCase, "TC"),
            title=title[:200],
            requirement_id=item.id,
            precondition=(request.form.get(f"rows-{idx}-precondition") or "").strip(),
            steps=(request.form.get(f"rows-{idx}-steps") or "").strip(),
            test_data=(request.form.get(f"rows-{idx}-test_data") or "").strip(),
            expected_result=(request.form.get(f"rows-{idx}-expected_result") or "").strip(),
            category=(request.form.get(f"rows-{idx}-category") or "功能测试").strip()[:64],
            priority=request.form.get(f"rows-{idx}-priority") or "medium",
            status=CASE_STATUS_CONFIRMED,
            source="llm",
            created_by_id=current_user.id,
        )
        if case.priority not in PRIORITY_NAMES:
            case.priority = "medium"
        db.session.add(case)
        db.session.flush()
        saved += 1

    db.session.commit()
    if saved:
        log_action("保存测试用例", target_type="requirement", target_id=item.id,
                   detail=f"保存 {saved} 条模型生成用例（需求 {item.code}）")
        flash(f"已保存 {saved} 条测试用例。", "success")
    else:
        flash("未选择任何有效用例，未保存。", "warning")
    if skipped:
        flash(f"有 {skipped} 条用例因名称为空被跳过，请补充后重新保存。", "warning")
    return redirect(url_for("testcases.list_testcases"))


# ---------------------------- R05 手工新增 ----------------------------
@bp.route("/new", methods=["GET", "POST"])
@login_required
def new():
    requirements_list = _active_requirements()
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        req_id = request.form.get("requirement_id", type=int)
        errors = flash_errors({"title": title}, [("title", "用例名称")])
        if not req_id:
            errors.append("必须关联一个软件需求。")
        if errors:
            for e in errors:
                flash(e, "danger")
            return render_template("testcases/form.html", item=None,
                                   requirements=requirements_list, form=request.form)

        case = TestCase(
            code=next_code(TestCase, "TC"),
            title=title[:200],
            requirement_id=req_id,
            precondition=(request.form.get("precondition") or "").strip(),
            steps=(request.form.get("steps") or "").strip(),
            test_data=(request.form.get("test_data") or "").strip(),
            expected_result=(request.form.get("expected_result") or "").strip(),
            category=(request.form.get("category") or "功能测试").strip()[:64],
            priority=request.form.get("priority") or "medium",
            status=CASE_STATUS_DRAFT,
            source="manual",
            created_by_id=current_user.id,
        )
        db.session.add(case)
        db.session.commit()
        log_action("新建测试用例", target_type="testcase", target_id=case.id,
                   detail=f"创建用例 {case.code}")
        flash(f"用例 {case.code} 创建成功。", "success")
        return redirect(url_for("testcases.detail", case_id=case.id))

    return render_template("testcases/form.html", item=None,
                           requirements=requirements_list, form={})


# ---------------------------- R04 查看 ----------------------------
@bp.route("/<int:case_id>")
@login_required
def detail(case_id):
    case = db.session.get(TestCase, case_id)
    if case is None or case.is_deleted:
        flash("测试用例不存在或已被删除。", "warning")
        return redirect(url_for("testcases.list_testcases"))
    return render_template("testcases/detail.html", item=case)


# ---------------------------- R04 编辑 ----------------------------
@bp.route("/<int:case_id>/edit", methods=["GET", "POST"])
@login_required
def edit(case_id):
    case = db.session.get(TestCase, case_id)
    if case is None or case.is_deleted:
        flash("测试用例不存在或已被删除。", "warning")
        return redirect(url_for("testcases.list_testcases"))
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        errors = flash_errors({"title": title}, [("title", "用例名称")])
        if errors:
            for e in errors:
                flash(e, "danger")
            return render_template("testcases/form.html", item=case,
                                   requirements=requirements_list, form=request.form)

        case.title = title[:200]
        case.requirement_id = request.form.get("requirement_id", type=int) or case.requirement_id
        case.precondition = (request.form.get("precondition") or "").strip()
        case.steps = (request.form.get("steps") or "").strip()
        case.test_data = (request.form.get("test_data") or "").strip()
        case.expected_result = (request.form.get("expected_result") or "").strip()
        case.category = (request.form.get("category") or "功能测试").strip()[:64]
        case.priority = request.form.get("priority") or "medium"
        db.session.commit()
        log_action("修改测试用例", target_type="testcase", target_id=case.id,
                   detail=f"更新用例 {case.code}")
        flash("测试用例已保存。", "success")
        return redirect(url_for("testcases.detail", case_id=case.id))

    return render_template("testcases/form.html", item=case,
                           requirements=requirements_list, form={
                               "title": case.title,
                               "requirement_id": case.requirement_id,
                               "precondition": case.precondition,
                               "steps": case.steps,
                               "test_data": case.test_data,
                               "expected_result": case.expected_result,
                               "category": case.category,
                               "priority": case.priority,
                           })


@bp.route("/<int:case_id>/confirm", methods=["POST"])
@login_required
def confirm(case_id):
    case = db.session.get(TestCase, case_id)
    if case is None or case.is_deleted:
        flash("测试用例不存在或已被删除。", "warning")
        return redirect(url_for("testcases.list_testcases"))
    case.status = CASE_STATUS_CONFIRMED
    db.session.commit()
    log_action("确认测试用例", target_type="testcase", target_id=case.id,
               detail=f"确认用例 {case.code}")
    flash(f"用例 {case.code} 已确认为正式用例。", "success")
    return redirect(url_for("testcases.detail", case_id=case.id))


# ---------------------------- R07 删除 ----------------------------
@bp.route("/<int:case_id>/delete", methods=["POST"])
@login_required
def delete(case_id):
    case = db.session.get(TestCase, case_id)
    if case is None or case.is_deleted:
        flash("测试用例不存在或已被删除。", "warning")
        return redirect(url_for("testcases.list_testcases"))

    if not current_user.is_admin and case.created_by_id != current_user.id:
        log_action("删除测试用例", target_type="testcase", target_id=case_id,
                   detail="权限不足：非本人创建且非管理员", result="failed")
        flash("只能删除本人创建的用例，或联系管理员。", "danger")
        return redirect(url_for("testcases.detail", case_id=case_id))

    case.is_deleted = True
    case.deleted_at = datetime.now()
    case.deleted_by_id = current_user.id
    db.session.commit()
    log_action("删除测试用例", target_type="testcase", target_id=case_id,
               detail=f"逻辑删除用例 {case.code}")
    flash(f"用例 {case.code} 已删除，可在回收站中还原。", "success")
    return redirect(url_for("testcases.list_testcases"))


@bp.route("/batch-delete", methods=["POST"])
@login_required
def batch_delete():
    """批量删除（仍为逻辑删除，可在回收站还原）。"""
    ids = request.form.getlist("ids")
    if not ids:
        flash("请先勾选需要删除的测试用例。", "warning")
        return redirect(url_for("testcases.list_testcases", **request.args.to_dict()))

    ok, denied = 0, 0
    for raw in ids:
        if not str(raw).isdigit():
            continue
        case = db.session.get(TestCase, int(raw))
        if case is None or case.is_deleted:
            continue
        if not current_user.is_admin and case.created_by_id != current_user.id:
            denied += 1
            continue
        case.is_deleted = True
        case.deleted_at = datetime.now()
        case.deleted_by_id = current_user.id
        ok += 1
    db.session.commit()

    if ok:
        log_action("批量删除测试用例", target_type="testcase", target_id="-",
                   detail=f"逻辑删除 {ok} 条用例")
        flash(f"已删除 {ok} 条测试用例，可在回收站中还原。", "success")
    if denied:
        log_action("批量删除测试用例", target_type="testcase", target_id="-",
                   detail=f"{denied} 条因权限不足被跳过", result="failed")
        flash(f"其中 {denied} 条不是本人创建，已跳过。", "warning")
    if not ok and not denied:
        flash("没有可删除的记录。", "warning")
    return redirect(url_for("testcases.list_testcases"))


# ---------------------------- R08 导出 ----------------------------
@bp.route("/export")
@login_required
def export():
    """R08 导出：支持 xlsx / csv / json / md，可导出当前筛选结果或勾选的用例。

    说明书 3.8 约束"具体格式待确定"，此处提供 4 种常用格式；
    导出内容始终来自数据库当前数据，与界面显示一致。
    """
    fmt = (request.args.get("format") or "xlsx").lower()
    f = _filters()
    ids = [int(i) for i in (request.args.get("ids") or "").split(",") if i.strip().isdigit()]

    query = db.session.query(TestCase).filter_by(is_deleted=False)
    if ids:
        query = query.filter(TestCase.id.in_(ids))
    else:
        query = _apply_filters(query, f)
    cases = query.order_by(TestCase.code.asc()).all()

    if not cases:
        flash("没有符合条件的测试用例可供导出。", "warning")
        return redirect(url_for("testcases.list_testcases", **request.args.to_dict()))

    headers = ["用例编号", "用例名称", "所属需求", "分类", "优先级", "状态",
               "前置条件", "测试步骤", "测试数据", "预期结果", "创建人", "更新时间"]

    def row_of(c):
        req = db.session.get(Requirement, c.requirement_id) if c.requirement_id else None
        return [
            c.code, c.title, f"{req.code} {req.name}" if req else "",
            c.category, c.priority_name, c.status_name,
            c.precondition, c.steps, c.test_data, c.expected_result,
            (c.creator.username if c.creator else ""),
            c.updated_at.strftime("%Y-%m-%d %H:%M") if c.updated_at else "",
        ]

    stamp = datetime.now().strftime("%Y%m%d%H%M%S")

    if fmt == "json":
        payload = [dict(zip(headers, row_of(c))) for c in cases]
        buf = io.BytesIO(json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"))
        mimetype, ext = "application/json", "json"
    elif fmt == "csv":
        text = io.StringIO()
        writer = csv.writer(text)
        writer.writerow(headers)
        for c in cases:
            writer.writerow(row_of(c))
        # 带 BOM，便于 Excel 正确识别中文
        buf = io.BytesIO(("\ufeff" + text.getvalue()).encode("utf-8"))
        mimetype, ext = "text/csv", "csv"
    elif fmt in ("md", "markdown"):
        lines = [f"# 测试用例清单", "",
                 f"导出时间：{datetime.now():%Y-%m-%d %H:%M}　用例数量：{len(cases)}", ""]
        for c in cases:
            req = db.session.get(Requirement, c.requirement_id) if c.requirement_id else None
            lines += [
                f"## {c.code} {c.title}", "",
                f"- 所属需求：{f'{req.code} {req.name}' if req else '—'}",
                f"- 分类：{c.category}　优先级：{c.priority_name}　状态：{c.status_name}", "",
                f"**前置条件**：{c.precondition or '—'}", "",
                "**测试步骤**：", "",
            ]
            steps = c.step_list or ["—"]
            for i, s in enumerate(steps, start=1):
                lines.append(f"{i}. {s}")
            lines += ["",
                      f"**测试数据**：{c.test_data or '—'}", "",
                      f"**预期结果**：{c.expected_result or '—'}", "",
                      f"**创建人**：{c.creator.username if c.creator else '—'}"
                      f"　**更新时间**：{c.updated_at.strftime('%Y-%m-%d %H:%M') if c.updated_at else '—'}",
                      "", "---", ""]
        buf = io.BytesIO("\n".join(lines).encode("utf-8"))
        mimetype, ext = "text/markdown", "md"
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "测试用例"
        ws.append(headers)
        head_fill = PatternFill("solid", fgColor="2F5597")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in cases:
            ws.append(row_of(c))
        widths = [12, 26, 24, 12, 8, 8, 20, 32, 16, 32, 10, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            for cell in r:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        mimetype, ext = ("application/vnd.openxmlformats-officedocument."
                         "spreadsheetml.sheet"), "xlsx"
    else:
        flash("不支持的导出格式。", "danger")
        return redirect(url_for("testcases.list_testcases"))

    buf.seek(0)
    log_action("导出测试用例", target_type="testcase", target_id="-",
               detail=f"导出 {len(cases)} 条用例，格式 {ext}")
    return send_file(buf, mimetype=mimetype, as_attachment=True,
                     download_name=f"testcases_{stamp}.{ext}")
